from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

BASE    = r'C:\VuzUC\MAGA\Nstya_SECOND'
IMG_DIR = os.path.join(BASE, 'L3_pages')

doc = Document()

sec = doc.sections[0]
sec.left_margin   = Cm(3)
sec.right_margin  = Cm(1.5)
sec.top_margin    = Cm(2)
sec.bottom_margin = Cm(2)


def set_spacing(p, line=1.5, sb=0, sa=0):
    pPr = p._p.get_or_add_pPr()
    spc = OxmlElement('w:spacing')
    spc.set(qn('w:line'),     str(int(line * 240)))
    spc.set(qn('w:lineRule'), 'auto')
    spc.set(qn('w:before'),   str(int(sb * 20)))
    spc.set(qn('w:after'),    str(int(sa * 20)))
    pPr.append(spc)


def add_run(p, text, bold=False):
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'),    'Times New Roman')
    ex = rPr.find(qn('w:rFonts'))
    if ex is not None:
        rPr.remove(ex)
    rPr.insert(0, rFonts)
    return run


def new_para(indent=True, center=False):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center' if center else 'both')
    pPr.append(jc)
    ind = OxmlElement('w:ind')
    if indent:
        ind.set(qn('w:firstLine'), '709')
    pPr.append(ind)
    set_spacing(p)
    return p


def bullet(text):
    p = new_para(indent=False)
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:left'),    '709')
    ind.set(qn('w:hanging'), '360')
    add_run(p, f'–\t{text}')
    return p


# ── ТЕКСТ ─────────────────────────────────────────────────────────────────────

p = new_para()
add_run(p, 'Цель работы: ', bold=True)
add_run(p, 'создать сборочные чертежи в NanoCAD, соблюдая ГОСТ 2.109-73. '
           'Соблюсти все размеры, требования к оформлению сборочного чертежа, '
           'технические требования, требования к спецификации в соответствии '
           'с ГОСТ 2.108-68, а также допуски и элементы сварки.')

p = new_para()
add_run(p, 'Ход работы:', bold=True)

p = new_para()
add_run(p, 'В ходе работы был создан сборочный чертёж в соответствии с требованиями ГОСТ, '
           'который содержит следующие элементы: изображение сборочной единицы, размеры, '
           'номера позиций, технические требования.')

p = new_para()
add_run(p, 'Все размеры, указанные в сборочном чертеже, разделены по следующим категориям '
           '(при необходимости): габаритные, установленные и присоединительные, '
           'эксплуатационные и монтажные.')

p = new_para()
add_run(p, 'Общие требования к сборочным чертежам, которые были учтены при выполнении задания:')

for b in [
    'все поверхности сопрягаемых деталей в месте соединения необходимо обозначить одной контурной линией;',
    'смежные детали в разрезах и сечениях обозначают штриховкой в различных направлениях;',
    'если число деталей больше двух — направление и частота штриховки меняется;',
    'сплошные детали в секущей плоскости вдоль оси или вдоль длинной стороны не штрихуются;',
    'собранные в сборочную единицу отдельные детали изображаются в рабочем положении;',
    'цельные, непустотелые детали изображаются в продольных разрезах незаштрихованными;',
    'для упрощения чтения сборочных чертежей на них допускается помещать изображение пограничных элементов и их размеры;',
    'линии невидимого контура применяют только для изображения простых элементов, '
    'когда выполнение разреза не упрощает чтение чертежа, а увеличивает его трудоёмкость;',
    'на сборочных чертежах перемещающиеся в работе части допускается изображать в крайнем '
    'или промежуточном положении с указанием соответствующего размера;',
    'клапанные устройства принято изображать закрытыми.',
]:
    bullet(b)

p = new_para()
add_run(p, 'Для изображений сборочных единиц существуют отдельные требования оформления, '
           'которые также были учтены при выполнении данной работы. К ним относятся:')

for b in [
    'количество изображений должно быть минимальным, но раскрывающим информацию '
    'о форме, расположении и характере соединения деталей;',
    'при разрезах все смежные детали необходимо заштриховать в противоположные стороны;',
    'на всех изображениях одна и та же деталь должна иметь одинаковую штриховку;',
    'на разрезах стандартные изделия не заштриховываются.',
]:
    bullet(b)

p = new_para()
add_run(p, 'К оформлению размеров на сборочных чертежах предъявляются следующие требования:')

for b in [
    'указываются габаритные размеры изделия, определяющие его внешние очертания;',
    'указываются установочные и присоединительные размеры, характеризующие установку '
    'изделия на месте его монтажа или присоединения к другому изделию;',
    'указываются другие необходимые и справочные размеры;',
    'все справочные размеры на чертеже отмечают знаком «*», а в технических требованиях '
    'выполняется надпись «*Размеры для справок».',
]:
    bullet(b)

p = new_para()
add_run(p, 'Помимо всего прочего, были учтены технические требования, включающие: '
           'подтверждение материала сертификатом при входном контроле, обеспечение '
           'идентификации материала на всех этапах технологического процесса, '
           'а также требования к сварным швам по ГОСТ 14771-76.')

p = new_para()
add_run(p, 'В ходе работы выполнены требования к спецификации, которые были освещены '
           'в презентации преподавателем (разделы, графы, позиционирование). '
           'В спецификацию входят разделы: документация, детали, стандартные изделия.')

p = new_para()
add_run(p, 'Также были учтены допуски и особенности сварки. '
           'Сварные швы выполнены по ГОСТ 14771-76 (сварка в защитных газах). '
           'Условные обозначения сварных швов проставлены на сборочных чертежах '
           'в соответствии с ГОСТ 2.312-72.')

p = new_para()
add_run(p, 'Таким образом, учитывая все особенности и нюансы работы, были выполнены '
           'сборочные чертежи трёх элементов (Балка 2А.012.001 СБ, '
           'Подставка 1 2А.012.002 СБ, Подставка 2 2А.012.003 СБ) '
           'и спецификации к ним. Результат работы представлен ниже.')

p = new_para()
add_run(p, 'Вывод: ', bold=True)
add_run(p, 'в ходе выполнения данной практической работы были созданы сборочные чертежи '
           'в NanoCAD с соблюдением ГОСТ 2.109-73. Были соблюдены все размеры, требования '
           'к оформлению сборочного чертежа, технические требования, требования к спецификации '
           'в соответствии с ГОСТ 2.108-68, а также допуски и элементы сварки.')

# ── ЧЕРТЕЖИ + СПЕЦИФИКАЦИИ ────────────────────────────────────────────────────

pairs = [
    ('chert_01.png', 'spec_01.png'),
    ('chert_02.png', 'spec_02.png'),
    ('chert_03.png', 'spec_03.png'),
]

for chert_file, spec_file in pairs:
    # Ландшафтная секция — сборочный чертёж
    sec_land = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec_land.orientation  = WD_ORIENT.LANDSCAPE
    sec_land.page_width   = Cm(29.7)
    sec_land.page_height  = Cm(21)
    sec_land.left_margin  = Cm(1)
    sec_land.right_margin = Cm(1)
    sec_land.top_margin   = Cm(1)
    sec_land.bottom_margin= Cm(1)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, sb=0, sa=0)
    p.add_run().add_picture(os.path.join(IMG_DIR, chert_file), width=Cm(27.7))

    # Портретная секция — спецификация
    sec_port = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec_port.orientation  = WD_ORIENT.PORTRAIT
    sec_port.page_width   = Cm(21)
    sec_port.page_height  = Cm(29.7)
    sec_port.left_margin   = Cm(3)
    sec_port.right_margin  = Cm(1.5)
    sec_port.top_margin    = Cm(2)
    sec_port.bottom_margin = Cm(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, sb=0, sa=0)
    p.add_run().add_picture(os.path.join(IMG_DIR, spec_file), width=Cm(17))

# ── СОХРАНЕНИЕ ────────────────────────────────────────────────────────────────

out = os.path.join(BASE, 'отчёт_пр3.docx')
doc.save(out)
print(f'Saved: {out}')
