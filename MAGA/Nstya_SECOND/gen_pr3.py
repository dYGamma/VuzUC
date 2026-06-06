from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

sec = doc.sections[0]
sec.left_margin   = Cm(3)
sec.right_margin  = Cm(1.5)
sec.top_margin    = Cm(2)
sec.bottom_margin = Cm(2)

def set_spacing(p, line=1.5, sb=0, sa=0):
    pPr = p._p.get_or_add_pPr()
    spc = OxmlElement('w:spacing')
    spc.set(qn('w:line'),     str(int(line * 240)))
    spc.set(qn('w:lineRule'), 'auto')
    spc.set(qn('w:before'),   str(int(sb * 20)))
    spc.set(qn('w:after'),    str(int(sa * 20)))
    pPr.append(spc)

def add_run(p, text, bold=False):
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'),    'Times New Roman')
    ex = rPr.find(qn('w:rFonts'))
    if ex is not None:
        rPr.remove(ex)
    rPr.insert(0, rFonts)
    return run

def new_para(indent=True, center=False):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center' if center else 'both')
    pPr.append(jc)
    ind = OxmlElement('w:ind')
    if indent:
        ind.set(qn('w:firstLine'), '709')
    pPr.append(ind)
    set_spacing(p)
    return p

def bullet(text):
    p = new_para(indent=False)
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:left'),    '709')
    ind.set(qn('w:hanging'), '360')
    add_run(p, f'–\t{text}')
    return p

# ── Страница 1: текст ─────────────────────────────────────────────────────────

p = new_para()
add_run(p, 'Цель работы: ', bold=True)
add_run(p, 'создать сборочные чертежи в NanoCAD, соблюдая ГОСТ 2.109-73. '
           'Соблюсти все размеры, требования к оформлению сборочного чертежа, '
           'технические требования, требования к спецификации в соответствии '
           'с ГОСТ 2.108-68, а также допуски и элементы сварки.')

p = new_para()
add_run(p, 'Ход работы:', bold=True)

p = new_para()
add_run(p, 'В ходе работы был создан сборочный чертёж в соответствии с требованиями ГОСТ, '
           'который содержит следующие элементы: изображение сборочной единицы, размеры, '
           'номера позиций, технические требования.')

p = new_para()
add_run(p, 'Все размеры, указанные в сборочном чертеже, разделены по следующим категориям: '
           'габаритные, установленные и присоединительные, эксплуатационные и монтажные.')

p = new_para()
add_run(p, 'При выполнении задания учтены следующие общие требования к сборочным чертежам:')

bullets = [
    'все поверхности сопрягаемых деталей в месте соединения обозначены одной контурной линией;',
    'смежные детали в разрезах и сечениях обозначены штриховкой в различных направлениях;',
    'если число деталей больше двух — направление и частота штриховки меняется;',
    'сплошные детали в секущей плоскости вдоль оси или вдоль длинной стороны не штрихуются;',
    'собранные в сборочную единицу детали изображаются в рабочем положении;',
    'для изображений сборочных единиц применяется минимальное количество изображений, '
     'достаточное для раскрытия информации о форме, расположении и характере соединения деталей;',
    'для каждой сборочной единицы составлена спецификация по ГОСТ 2.108-68.',
]
for b in bullets:
    bullet(b)

p = new_para()
add_run(p, 'Вывод: ', bold=True)
add_run(p, 'таким образом, учитывая все особенности и нюансы работы, были выполнены '
           'сборочные чертежи трёх элементов (Балка, Подставка 1, Подставка 2) '
           'и спецификации к ним в соответствии с требованиями ГОСТ 2.109-73 и ГОСТ 2.108-68.')

# ── Страницы 2-7: чертежи и спецификации ────────────────────────────────────

IMG_DIR = r'C:\VuzUC\MAGA\Nstya_SECOND\L3_pages'

pages = [
    'chert_01.png',  # Балка СБ
    'chert_02.png',  # Подставка 1 СБ
    'chert_03.png',  # Подставка 2 СБ
    'spec_01.png',   # Балка спецификация
    'spec_02.png',   # Подставка 1 спецификация
    'spec_03.png',   # Подставка 2 спецификация
]

for fname in pages:
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, sb=0, sa=0)
    run = p.add_run()
    run.add_picture(os.path.join(IMG_DIR, fname), width=Cm(17))

out = r'C:\VuzUC\MAGA\Nstya_SECOND\отчёт_пр3.docx'
doc.save(out)
print(f'Saved: {out}')

d2 = Document(out)
imgs = sum(1 for p in d2.paragraphs if p._p.find('.//' + qn('a:blip')) is not None)
print(f'Images: {imgs}')
