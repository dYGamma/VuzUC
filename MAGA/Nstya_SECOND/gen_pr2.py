from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

sec = doc.sections[0]
sec.left_margin   = Cm(3)
sec.right_margin  = Cm(1.5)
sec.top_margin    = Cm(2)
sec.bottom_margin = Cm(2)

def set_spacing(p, line=1.5, sb=0, sa=0):
    pPr = p._p.get_or_add_pPr()
    spc = OxmlElement('w:spacing')
    spc.set(qn('w:line'),     str(int(line * 240)))
    spc.set(qn('w:lineRule'), 'auto')
    spc.set(qn('w:before'),   str(int(sb * 20)))
    spc.set(qn('w:after'),    str(int(sa * 20)))
    pPr.append(spc)

def add_run(p, text, bold=False):
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'),    'Times New Roman')
    ex = rPr.find(qn('w:rFonts'))
    if ex is not None:
        rPr.remove(ex)
    rPr.insert(0, rFonts)
    return run

def new_para(indent=True, center=False):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center' if center else 'both')
    pPr.append(jc)
    ind = OxmlElement('w:ind')
    if indent:
        ind.set(qn('w:firstLine'), '709')
    pPr.append(ind)
    set_spacing(p)
    return p

def numbered_item(n, text):
    p = new_para(indent=False)
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:left'),    '709')
    ind.set(qn('w:hanging'), '709')
    add_run(p, f'{n}.\t{text}')
    return p

# ── Страница 1: текст ────────────────────────────────────────────────────────

p = new_para()
add_run(p, 'Цель работы: ', bold=True)
add_run(p, 'создать чертёж из детали и перенести в NanoCAD, где оформить '
           'в соответствии с ГОСТ все проекции.')

p = new_para()
add_run(p, 'Ход работы:', bold=True)

steps = [
    'Перенести все проекции в шаблон, во вкладку «Модель».',
    'Для удобства разбить область «Модели» по узлам.',
    'Выделить все основные линии и установить вес линии в 0.40.',
    'Расставить осевые весом 0.18 мм (штрихпунктирные).',
    'Переходим в область листа. С помощью «видового экрана» подбираем необходимый масштаб.',
    'Возвращаемся в «Модель».',
    'Наносим размеры, соблюдая ГОСТ и выбранный масштаб '
    '(шрифты Stand1, Stand2, Stand4, Stand5, Stand10, Stand20 соответствуют '
    'масштабу 1:1, 1:2, 1:4, 1:5, 1:10, 1:20).',
    'Указываем отклонения А-А.',
    'Указываем необходимую обработку (шероховатость).',
    'Переходим в «Лист».',
    'Выравниваем, указываем общую шероховатость.',
    'Заполняем технические требования.',
    'Заполняем рамку.',
]
for i, s in enumerate(steps, 1):
    numbered_item(i, s)

p = new_para()
add_run(p, 'Вывод: ', bold=True)
add_run(p, 'в данной практической работе были выполнены чертежи каждой '
           'отдельной детали в ПО NanoCAD, в соответствии с ГОСТ.')

# ── Страницы 2-21: чертежи ───────────────────────────────────────────────────

IMG_DIR = r'C:\VuzUC\MAGA\Nstya_SECOND\L2_pages'

for page_num in range(1, 21):
    # page break before each drawing
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, sb=0, sa=0)
    run = p.add_run()
    path = os.path.join(IMG_DIR, f'page_{page_num:02d}.png')
    run.add_picture(path, width=Cm(17))

out = r'C:\VuzUC\MAGA\Nstya_SECOND\отчёт_пр2_v2.docx'
doc.save(out)
print(f'Saved: {out}')

d2 = Document(out)
imgs = sum(1 for p in d2.paragraphs if p._p.find('.//' + qn('a:blip')) is not None)
print(f'Images: {imgs}')
